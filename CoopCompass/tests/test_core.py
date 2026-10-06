"""Unit tests"""
from datetime import date, datetime, timedelta

import pytest

from ai import grounding
from backend.app.services import analytics, job_parser, matching, metrics, resume_parser, skills
from backend.app.services.skills import extract_skills
from tests.conftest import SAMPLE_JD, SAMPLE_RESUME


def test_skill_extraction_handles_ambiguous_names():
    text = "Python, SQL and R. We express interest; the rest of spring 2024. Go to market. Excel at math. React Native app."
    got = extract_skills(text)
    assert {"Python", "SQL", "R"} <= set(got)
    assert "Go" not in got and "Express" not in got and "Excel" not in got and "React" not in got
    assert "React Native" in got and "Go-to-Market" in got
    assert extract_skills("Built dashboards in Power BI and Tableau with PostgreSQL") == ["Data Visualization", "Power BI", "Tableau", "PostgreSQL"] \
        or set(extract_skills("Built dashboards in Power BI and Tableau with PostgreSQL")) >= {"Power BI", "Tableau", "PostgreSQL"}


def test_skill_boundaries_do_not_match_substrings():
    assert "Java" not in extract_skills("JavaScript developer")
    assert "SQL" not in extract_skills("MongoDB and NoSQL stores")
    assert skills.canonicalize("postgres") == "PostgreSQL"


def test_related_requires_same_specific_family():
    assert skills.related("Azure", "AWS")
    assert not skills.related("Python", "Java")
    assert not skills.related("AWS", "AWS")


def test_resume_parser_sections_dates_education():
    d = resume_parser.parse_resume(SAMPLE_RESUME)
    assert {"Python", "SQL", "Power BI", "FastAPI", "PostgreSQL"} <= set(d["skills"])
    sections = {e["section"] for e in d["evidence_items"]}
    assert {"Experience", "Projects", "Skills", "Education"} <= sections
    exp = [e for e in d["evidence_items"] if e["section"] == "Experience"]
    assert any("Python" in e["skills"] and "SQL" in e["skills"] for e in exp)
    assert d["education"][0]["degree"] == "Master of Science" and d["education"][0]["status"] == "in_progress"
    assert d["education"][1]["degree"] == "B.Tech"
    assert d["experience_months"] == 27


def test_experience_months_merges_overlaps_and_present():
    lines = ["Intern  Jan 2023 - Jun 2023", "Part-time  Mar 2023 - Aug 2023", "Role  Jan 2025 - Present"]
    assert resume_parser.experience_months(lines, today=date(2025, 3, 15)) == 8 + 3
    assert resume_parser.experience_months(["no dates here"]) == 0


def test_resume_text_extraction_txt_and_rejects_unknown():
    text, _ = resume_parser.extract_text("r.txt", b"hello world")
    assert text == "hello world"
    with pytest.raises(ValueError):
        resume_parser.extract_text("r.exe", b"x")


def test_docx_extraction(tmp_path):
    import docx

    d = docx.Document()
    d.add_paragraph("Experience")
    d.add_paragraph("Built SQL reports for finance operations team")
    p = tmp_path / "r.docx"
    d.save(p)
    text, _ = resume_parser.extract_text("r.docx", p.read_bytes())
    assert "SQL reports" in text


def test_job_parser_sections_and_facts():
    j = job_parser.parse_job(SAMPLE_JD, title="Data Analyst Co-op", company="Example Co")
    assert j["required_skills"][:2] == ["Python", "SQL"] and "A/B Testing" in j["required_skills"]
    assert j["preferred_skills"][:2] == ["AWS", "Azure"] and "Tableau" in j["preferred_skills"]
    assert j["minimum_experience_years"] == 1.0
    assert j["education"] == "bachelor"
    assert j["employment_type"] == "co-op"
    assert j["location"] == "Boston, MA"
    assert len(j["work_authorization_notes"]) == 1 and "sponsor" in j["work_authorization_notes"][0]


def test_job_parser_does_not_invent_missing_facts():
    j = job_parser.parse_job("We are hiring a thoughtful person to help with various tasks across the team.", title="Helper", company="X")
    assert j["minimum_experience_years"] is None and j["education"] is None and j["application_deadline"] is None
    assert j["salary_text"] is None and j["work_authorization_notes"] == []


def test_job_parser_flags_unsectioned_postings():
    j = job_parser.parse_job("Looking for someone strong in Python and SQL to join our analytics group right away.", "A", "B")
    assert j["required_skills"] == ["Python", "SQL"]
    assert any("treated as required" in n for n in j["parse_notes"])


def test_salary_and_deadline_parsing():
    assert job_parser.parse_salary("Pay: $25 - $32 per hour")[1:] == (25.0, 32.0)
    assert job_parser.parse_salary("$80,000 - $95,000 / year")[1:] == (80000.0, 95000.0)
    assert job_parser.parse_deadline("Apply by March 15, 2027.") == datetime(2027, 3, 15)
    assert job_parser.parse_deadline("nothing") is None


def _profile(**kw):
    p = {"target_roles": ["Data Analyst"], "skills": ["Python", "SQL"], "experience_months": 24,
         "education": [{"degree": "Master of Science", "field": "Analytics", "school": "X", "status": "in_progress", "year": ""}],
         "preferred_locations": ["Boston, MA"], "remote_preference": "any", "employment_types": ["co-op"],
         "industries": [], "requires_sponsorship": None, "preferred_technologies": [], "salary_min": None, "headline": "",
         "evidence_items": [
             {"id": "e1", "section": "Experience", "text": "Automated reporting with Python and SQL", "skills": ["Python", "SQL"]},
             {"id": "e2", "section": "Skills", "text": "Python, SQL, Tableau", "skills": ["Python", "SQL", "Tableau"]}],
         "match_weights": matching.DEFAULT_WEIGHTS}
    p.update(kw)
    return p


def _job(**kw):
    j = {"id": 1, "title": "Data Analyst Co-op", "company": "Co", "location": "Boston, MA", "employment_type": "co-op",
         "required_skills": ["Python", "SQL"], "preferred_skills": [], "minimum_experience_years": 1.0, "education": "bachelor",
         "responsibilities": ["Analyze product usage data with SQL"], "work_authorization_notes": [], "is_remote": False,
         "salary_min": None, "salary_max": None, "application_deadline": None, "raw_description": "Data analyst role"}
    j.update(kw)
    return j


def test_weights_normalise_to_100():
    w = matching.normalize_weights({"skills": 70, "experience": 0, "role": 0, "education": 0, "location": 0, "preferences": 0})
    assert abs(sum(w.values()) - 100) < 1e-6 and w["skills"] == 100
    assert abs(sum(matching.normalize_weights(None).values()) - 100) < 1e-6
    assert matching.parse_weights_param("skills:50,role:50") == {"skills": 50.0, "role": 50.0}


def test_perfect_match_scores_high_and_overall_is_sum_of_points():
    m = matching.compute_match(_job(), _profile())
    assert m["overall"] >= 90
    assert m["overall"] == pytest.approx(sum(b["points"] for b in m["breakdown"]), abs=0.11)
    for b in m["breakdown"]:
        assert 0 <= b["points"] <= b["weight"] + 0.05
    assert not m["gaps"] and {s["skill"] for s in m["strong"]} == {"Python", "SQL"}


def test_skill_levels_demonstrated_listed_related_gap():
    prof = _profile(skills=["Python", "SQL", "Tableau", "Azure"], evidence_items=[
        {"id": "e1", "section": "Experience", "text": "Built reports with Python", "skills": ["Python"]},
        {"id": "e2", "section": "Experience", "text": "Deployed services on Azure", "skills": ["Azure"]}])
    job = _job(required_skills=["Python", "SQL", "AWS", "Kubernetes"])
    m = matching.compute_match(job, prof)
    strong = {s["skill"] for s in m["strong"]}
    partial = {p["skill"]: p["reason"] for p in m["partial"]}
    gaps = {g["requirement"] for g in m["gaps"]}
    assert strong == {"Python"}
    assert "listed in your skills" in partial["SQL"].lower()
    assert "Related" in partial["AWS"] and "Azure" in partial["AWS"]
    assert gaps == {"Kubernetes"}


def test_experience_and_education_scoring():
    m = matching.compute_match(_job(minimum_experience_years=4.0), _profile(experience_months=24))
    exp = next(b for b in m["breakdown"] if b["key"] == "experience")
    assert exp["fraction"] == pytest.approx(0.5)
    assert any(g["kind"] == "experience" for g in m["gaps"])
    m2 = matching.compute_match(_job(education="phd"), _profile())
    edu = next(b for b in m2["breakdown"] if b["key"] == "education")
    assert edu["fraction"] == 0.5
    m3 = matching.compute_match(_job(minimum_experience_years=None), _profile())
    assert next(b for b in m3["breakdown"] if b["key"] == "experience")["fraction"] == 0.8


def test_location_and_sponsorship_logic():
    m = matching.compute_match(_job(location="Seattle, WA"), _profile())
    loc = next(b for b in m["breakdown"] if b["key"] == "location")
    assert loc["fraction"] == 0.3 and any(c["type"] == "location" for c in m["concerns"])
    remote = matching.compute_match(_job(is_remote=True, location=None), _profile())
    assert next(b for b in remote["breakdown"] if b["key"] == "location")["fraction"] == 1.0
    notes = ["We are unable to sponsor work visas."]
    base = matching.compute_match(_job(work_authorization_notes=notes), _profile(requires_sponsorship=False))
    needs = matching.compute_match(_job(work_authorization_notes=notes), _profile(requires_sponsorship=True))
    pb = lambda m: next(b for b in m["breakdown"] if b["key"] == "preferences")["points"]
    assert pb(needs) < pb(base)
    assert any("exclude sponsorship" in c["message"] for c in needs["concerns"])


def test_custom_weights_change_overall_but_not_dimension_fractions():
    a = matching.compute_match(_job(location="Seattle, WA"), _profile())
    b = matching.compute_match(_job(location="Seattle, WA"), _profile(), weights={"skills": 10, "location": 90, "experience": 0, "role": 0, "education": 0, "preferences": 0})
    assert b["overall"] < a["overall"]
    assert [x["fraction"] for x in a["breakdown"]] == [x["fraction"] for x in b["breakdown"]]


def test_no_profile_evidence_never_invents_matches():
    prof = _profile(skills=[], evidence_items=[])
    m = matching.compute_match(_job(), prof)
    assert not m["strong"] and {g["requirement"] for g in m["gaps"]} == {"Python", "SQL"}
    assert all(g["message"] == "No evidence found in current resume." for g in m["gaps"] if g["kind"].startswith("skill"))


def test_gap_analysis_responsibility_coverage():
    g = matching.gap_analysis(_job(responsibilities=["Analyze product usage data with SQL and Python", "Manage a fleet of forklifts in a warehouse"]), _profile())
    by = {r["text"][:7]: r for r in g["responsibilities"]}
    assert by["Analyze"]["status"] in {"present", "weak"} and by["Analyze"]["best_evidence"]["id"] == "e1"
    assert by["Manage "]["status"] == "gap" and by["Manage "]["best_evidence"] is None


def test_grounding_rejects_invented_skills_and_numbers():
    src = "Built dashboards in Power BI for the operations team"
    assert grounding.check_rewrite(src, "Built Power BI dashboards that gave the operations team clear visibility")[0]
    ok, why = grounding.check_rewrite(src, "Built Power BI and Tableau dashboards for operations")
    assert not ok and "Tableau" in why
    ok, why = grounding.check_rewrite(src, "Built Power BI dashboards that cut reporting time by 40%")
    assert not ok and "40%" in why
    assert not grounding.check_rewrite(src, "  ")[0]


def test_ai_service_discards_ungrounded_gemini_rewrite(db, monkeypatch):
    import json

    from ai import provider
    from backend.app.services import ai_service

    monkeypatch.setattr(provider, "gemini_configured", lambda: True)
    prof = _profile()
    eid = prof["evidence_items"][0]["id"]
    calls = []

    def fake(prompt, json_mode=False, temperature=0.3):
        calls.append(prompt)
        return json.dumps([{"id": eid, "rewrite": "Automated reporting with Python, SQL and Kubernetes"}])

    monkeypatch.setattr(provider, "generate", fake)
    r = ai_service.tailor(db, _job(), prof)
    assert r["provider"] == "gemini" and r["suggestions"][0]["suggested"] is None
    assert "grounding check" in r["suggestions"][0]["rationale"]
    r2 = ai_service.tailor(db, _job(), prof)
    assert r2["cached"] is True and len(calls) == 1
    assert ai_service.requests_today(db) == 1


def test_ai_daily_limit_blocks_calls(db, monkeypatch):
    from ai import provider
    from backend.app.services import ai_service

    monkeypatch.setattr(provider, "gemini_configured", lambda: True)
    monkeypatch.setenv("AI_DAILY_REQUEST_LIMIT", "0")
    monkeypatch.setattr(provider, "generate", lambda *a, **k: pytest.fail("API must not be called past the limit"))
    r = ai_service.tailor(db, _job(), _profile())
    assert r["provider"] == "local" and "limit" in r["notice"].lower()


def _seed_apps(db, specs):
    from backend.app.models import Job, TimeRecord
    from backend.app.services import store

    for i, (title, score, path, base, assisted) in enumerate(specs):
        job, _ = store.create_job(db, company=f"Co{i}", title=title, description=f"{title} role needing Python and SQL skills.")
        app = store.create_application(db, job, path[0], {"exists": False})
        app.match_score = score
        db.commit()
        for st in path[1:]:
            store.set_status(db, app, st)
        if base:
            app.baseline_minutes = base
        if assisted:
            db.add(TimeRecord(job_id=job.id, kind="application_preparation", seconds=assisted * 60))
        db.commit()


def test_analytics_withholds_rates_until_enough_history(db):
    _seed_apps(db, [("Data Analyst", 80, ["Preparing", "Applied", "Interview"], None, None)])
    a = analytics.compute_analytics(db, {}, {})
    assert a["sufficient_data"] is False and a["insufficient_message"] == "Not enough application history yet."
    assert a["conversion"]["interview_rate"] == {"numerator": 1, "denominator": 1, "rate": None}


def test_analytics_rates_with_enough_history(db):
    specs = [("Data Analyst", 85, ["Preparing", "Applied", "Assessment", "Interview", "Offer"], None, None),
             ("Data Analyst", 90, ["Preparing", "Applied", "Interview"], None, None),
             ("Software Engineer", 55, ["Preparing", "Applied", "Rejected"], None, None),
             ("Product Manager Intern", 40, ["Preparing", "Applied"], None, None),
             ("ML Engineer", 72, ["Preparing", "Applied", "Assessment", "Rejected"], None, None),
             ("Data Analyst", 60, ["Saved"], None, None)]
    _seed_apps(db, specs)
    a = analytics.compute_analytics(db, {}, {})
    c = a["conversion"]
    assert a["sufficient_data"] and c["applied"] == 5
    assert c["interview_rate"] == {"numerator": 2, "denominator": 5, "rate": 0.4}
    assert c["assessment_rate"]["numerator"] == 2
    assert c["interview_to_offer_rate"] == {"numerator": 1, "denominator": 2, "rate": 0.5}
    fit = a["fit_conversion"]
    assert fit["high_fit"]["applications"] == 3 and fit["high_fit"]["interviews"] == 2
    assert fit["low_fit"] == {"applications": 2, "interviews": 0, "rate": 0.0}
    roles = {r["role"]: r for r in a["role_performance"]}
    assert roles["Data / Analytics"]["applications"] == 2 and roles["Data / Analytics"]["interviews"] == 2
    assert a["summary"]["applications_submitted"] == 5 and a["summary"]["offers"] == 1
    assert a["learning"]["stage"] == "collecting"
    assert {s["stage"]: s["count"] for s in a["funnel"]}["interview_received"] == 2


def test_funnel_is_monotone_by_construction(db):
    _seed_apps(db, [("Data Analyst", 80, ["Preparing", "Applied"], None, None)])
    counts = [s["count"] for s in analytics.compute_funnel(db)]
    assert counts == sorted(counts, reverse=True)


def test_time_saved_calculation(db):
    _seed_apps(db, [("Data Analyst", 80, ["Preparing", "Applied"], 120, 30),
                    ("Data Analyst", 80, ["Preparing", "Applied"], 90, 30),
                    ("Data Analyst", 80, ["Preparing", "Applied"], 100, 20),
                    ("Data Analyst", 80, ["Preparing"], None, 15)])
    m = metrics.compute_product_metrics(db)
    assert m["observations_with_baseline"] == 3 and m["claimable"] is False
    assert m["median_baseline_minutes"] == 100 and m["median_assisted_minutes"] == 30
    assert m["median_time_saved_minutes"] == 80
    assert m["median_time_reduction_pct"] == 75.0
    assert m["applications_started"] == 4 and m["applications_completed"] == 3
    assert m["weekly_estimated_time_saved_minutes"] is not None
    assert m["real_users"] == 0 and m["research_participants"] == 0


def test_time_saved_is_null_without_observations(db):
    m = metrics.compute_product_metrics(db)
    for k in ["median_baseline_minutes", "median_assisted_minutes", "median_time_saved_minutes", "median_time_reduction_pct",
              "weekly_estimated_time_saved_minutes", "resume_suggestion_acceptance_rate", "application_to_interview_rate"]:
        assert m[k] is None
    assert m["observations_with_baseline"] == 0 and m["claimable"] is False


def test_research_metrics_empty_is_collection_in_progress():
    r = metrics.compute_research_metrics([])
    assert r["status"] == "collection_in_progress" and r["message"] == "User research collection in progress."
    assert r["respondent_count"] == 0 and r["median_weekly_search_hours"] is None


def test_research_metrics_from_csv():
    csv_text = ("respondent_id,weekly_search_hours,minutes_per_application,pain_points,tools,desired_features\n"
                "r1,6,45,Too many tabs;Tracking,LinkedIn;Spreadsheet,Fit ranking;Tracker\n"
                "r2,10,30,Tracking,LinkedIn;Handshake,Fit ranking\n"
                "r3,4,60,Tracking;Tailoring,Spreadsheet,Tracker\n")
    r = metrics.compute_research_metrics(metrics.read_survey_csv(csv_text))
    assert r["respondent_count"] == 3 and r["median_weekly_search_hours"] == 6.0
    assert r["mean_minutes_per_application"] == 45.0
    assert r["top_pain_points"][0] == {"item": "Tracking", "count": 3}
    assert {"item": "Fit ranking", "count": 2} in r["feature_preferences"]


PDF_STYLE_RESUME = (
    "Sam Example\nSUMMARY\nA short summary.\nEDUCATION\n"
    "Example State University Sep 2024 \x96 Dec 2099 \n"
    "Master of Science in Information Systems; GPA: 3.8/4.0 Boston, MA \n"
    "Sample Institute, Example University Aug 2018 \x96 Jun 2022 \n"
    "Bachelor of Engineering in Electronics and Telecommunication; CGPA: 3.5 Pune, India \n"
    "EXPERIENCE\nAcme Corp Jul 2025 \x96 Dec 2025 \nSoftware Engineer Co-op Austin, TX \n"
    "\x95 Built an ordering platform with FastAPI and PostgreSQL for 50+ users and cut manual entry by 60% \n"
    "and shortened the review cycle for the operations team \n"
    "\x95 Planned sprints with the client team and demoed each release to leadership \n"
    "PRODUCT PROJECTS\nDemo Tool  |  Python, React GitHub \n"
    "\x95 Built a tool in Python that clusters feedback into themes for product managers \n"
    "SKILLS\nPython, SQL\n"
)


def test_pdf_style_layout_wrapped_bullets_dates_and_labels():
    d = resume_parser.parse_resume(PDF_STYLE_RESUME)
    edu = d["education"]
    assert edu[0]["school"] == "Example State University" and edu[0]["status"] == "in_progress"
    assert edu[0]["field"] == "Information Systems"
    assert edu[1]["school"] == "Sample Institute, Example University" and edu[1]["status"] == "completed"
    exp = [e for e in d["evidence_items"] if e["section"] == "Experience"]
    assert len(exp) == 2
    assert exp[0]["text"].endswith("operations team") and "review cycle" in exp[0]["text"]
    assert exp[0]["context"].startswith("Acme Corp - Software Engineer Co-op")
    assert {"FastAPI", "PostgreSQL"} <= set(exp[0]["skills"])
    proj = [e for e in d["evidence_items"] if e["section"] == "Projects"]
    assert proj and proj[0]["context"] == "Demo Tool"
    assert d["experience_months"] == 6


def test_skills_listed_without_evidence_is_flagged_and_scores_lower():
    job = _job(required_skills=["Python", "SQL", "Tableau"], minimum_experience_years=None)
    listed_only = _profile(evidence_items=[], skills=["Python", "SQL", "Tableau"])
    backed = _profile(skills=["Python", "SQL", "Tableau"], evidence_items=[
        {"id": "e1", "section": "Experience", "text": "Built Tableau reports with Python and SQL", "skills": ["Python", "SQL", "Tableau"]}])
    weak, strong = matching.compute_match(job, listed_only), matching.compute_match(job, backed)
    assert weak["low_evidence"] is True and strong["low_evidence"] is False
    assert any(c["type"] == "low_evidence" for c in weak["concerns"])
    assert strong["overall"] - weak["overall"] >= 15
    skills_dim = next(b for b in weak["breakdown"] if b["key"] == "skills")
    assert skills_dim["fraction"] == pytest.approx(0.3)
