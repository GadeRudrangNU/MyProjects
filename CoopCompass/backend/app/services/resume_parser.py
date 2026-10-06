"""Resume text extraction and parsing"""
from __future__ import annotations

import io
import re
import uuid
from datetime import date

from .skills import extract_skills

SECTION_ALIASES = {
    "experience": ["experience", "work experience", "professional experience", "employment", "work history",
                   "relevant experience", "internship experience", "internships"],
    "education": ["education", "academic background", "academics"],
    "projects": ["projects", "product projects", "academic projects", "personal projects", "selected projects", "project experience"],
    "skills": ["skills", "technical skills", "core competencies", "technologies", "tools", "skills & tools",
               "skills and tools", "technical proficiencies"],
    "summary": ["summary", "profile", "objective", "professional summary", "about"],
    "other": ["certifications", "certificates", "awards", "publications", "leadership", "activities",
              "volunteering", "honors", "extracurricular", "extracurriculars", "interests"],
}
_HEADING_LOOKUP = {a: k for k, v in SECTION_ALIASES.items() for a in v}
_BULLET = re.compile(r"^\s*[•●▪‣\-\*–—·\x95]\s+")
_MONTHS = {m: i for i, m in enumerate(
    ["jan", "feb", "mar", "apr", "may", "jun", "jul", "aug", "sep", "oct", "nov", "dec"], start=1)}
_DATE_RANGE = re.compile(
    r"(?P<m1>[A-Za-z]{3,9})\.?\s+(?P<y1>(?:19|20)\d{2})\s*(?:-|–|—|\x96|to)\s*"
    r"(?:(?P<m2>[A-Za-z]{3,9})\.?\s+(?P<y2>(?:19|20)\d{2})|(?P<present>present|current|now|ongoing))",
    re.IGNORECASE,
)
_NUM_DATE_RANGE = re.compile(
    r"(?P<m1>\d{1,2})/(?P<y1>(?:19|20)\d{2})\s*(?:-|–|—|\x96|to)\s*"
    r"(?:(?P<m2>\d{1,2})/(?P<y2>(?:19|20)\d{2})|(?P<present>present|current|now))",
    re.IGNORECASE,
)
_DEGREE = re.compile(
    r"\b(Ph\.?\s?D\.?|Doctor(?:ate)?(?: of Philosophy)?|MBA|"
    r"(?:Master|Bachelor)(?:'s)?(?: of (?:Science|Arts|Engineering|Technology|Business Administration|Fine Arts|Applied Science))?|"
    r"M\.?S\.?c?|M\.?Eng\.?|B\.?S\.?c?|B\.?A\.?|B\.?Tech\.?|B\.?E\.?|Associate(?:'s)?)(?![A-Za-z])",
)
_SCHOOL = re.compile(r"\b(University|College|Institute|School of|Polytechnic|Academy)\b", re.IGNORECASE)
_FIELD = re.compile(r"\bin\s+([A-Z][A-Za-z&' ]{2,60}?)(?=\s*(?:,|;|\||\(|\d|$|\s-\s|\sat\s))")


def extract_text(filename: str, data: bytes) -> tuple[str, list[str]]:
    name = filename.lower()
    warnings: list[str] = []
    if name.endswith(".pdf"):
        from pypdf import PdfReader

        reader = PdfReader(io.BytesIO(data))
        text = "\n".join((p.extract_text() or "") for p in reader.pages)
        if len(text.strip()) < 50:
            warnings.append("Very little text could be extracted (scanned/image PDF?). Add your content to the profile manually.")
    elif name.endswith(".docx"):
        import docx

        d = docx.Document(io.BytesIO(data))
        parts = [p.text for p in d.paragraphs]
        for t in d.tables:
            for row in t.rows:
                parts.append(" | ".join(c.text for c in row.cells))
        text = "\n".join(parts)
    elif name.endswith((".txt", ".md")):
        text = data.decode("utf-8", errors="ignore")
    else:
        raise ValueError("Unsupported file type. Upload a PDF, DOCX or TXT file.")
    return text, warnings


def split_sections(text: str) -> dict[str, list[str]]:
    sections: dict[str, list[str]] = {"header": []}
    current = "header"
    for raw in text.splitlines():
        line = raw.strip()
        if not line:
            continue
        key = re.sub(r"[^a-z& ]", "", line.lower()).strip()
        if key in _HEADING_LOOKUP and len(line) < 40:
            current = _HEADING_LOOKUP[key]
            sections.setdefault(current, [])
            continue
        sections.setdefault(current, []).append(line)
    return sections


def _to_month_index(month: str, year: int) -> int | None:
    mi = _MONTHS.get(month[:3].lower())
    return None if mi is None else year * 12 + mi


def experience_months(lines: list[str], today: date | None = None) -> int:
    today = today or date.today()
    now_idx = today.year * 12 + today.month
    spans: list[tuple[int, int]] = []
    for line in lines:
        for m in _DATE_RANGE.finditer(line):
            start = _to_month_index(m.group("m1"), int(m.group("y1")))
            if start is None:
                continue
            if m.group("present"):
                end = now_idx
            else:
                end = _to_month_index(m.group("m2") or "", int(m.group("y2") or 0))
            if end is not None and end >= start:
                spans.append((start, end + 1))
        for m in _NUM_DATE_RANGE.finditer(line):
            start = int(m.group("y1")) * 12 + int(m.group("m1"))
            end = now_idx if m.group("present") else int(m.group("y2")) * 12 + int(m.group("m2"))
            if end >= start:
                spans.append((start, end + 1))
    if not spans:
        return 0
    spans.sort()
    total, (cs, ce) = 0, spans[0]
    for s, e in spans[1:]:
        if s <= ce:
            ce = max(ce, e)
        else:
            total += ce - cs
            cs, ce = s, e
    total += ce - cs
    return total


def _is_header_line(line: str) -> bool:
    return bool(_DATE_RANGE.search(line) or _NUM_DATE_RANGE.search(line)) and not _BULLET.match(line)


def _clean_header(line: str) -> str:
    return _DATE_RANGE.sub("", _NUM_DATE_RANGE.sub("", line)).strip(" |,-\u2013\u2014\x96")


def evidence_from_section(lines: list[str], section: str) -> tuple[list[dict], list[dict]]:
    bullets: list[tuple[str, str]] = []
    projects: list[dict] = []
    context = ""
    open_bullet = False
    expect_role = False
    for line in lines:
        if _BULLET.match(line):
            bullets.append((context, _BULLET.sub("", line).strip()))
            open_bullet, expect_role = True, False
        elif _is_header_line(line):
            context = _clean_header(line)
            open_bullet, expect_role = False, True
        elif open_bullet:
            ctx, txt = bullets[-1]
            bullets[-1] = (ctx, txt + " " + line.strip())
        elif expect_role:
            context = f"{context} - {line.strip()}"
            expect_role = False
        else:
            context = line.split("|")[0].strip()
            if section == "Projects" and context:
                projects.append({"name": context, "text": ""})
    items: list[dict] = []
    for ctx, text in bullets:
        if len(text) < 15:
            continue
        items.append({"id": uuid.uuid4().hex[:10], "section": section, "text": text,
                      "skills": extract_skills(text), "context": ctx})
        for p in projects:
            if p["name"] == ctx and not p["text"]:
                p["text"] = text
    return items, projects


def _ends_in_future(line: str, today: date | None = None) -> bool:
    today = today or date.today()
    m = _DATE_RANGE.search(line)
    if not m or m.group("present"):
        return bool(m)
    end = _to_month_index(m.group("m2") or "", int(m.group("y2") or 0))
    return end is not None and end > today.year * 12 + today.month


def parse_education(lines: list[str]) -> list[dict]:
    entries: list[dict] = []
    cur: dict | None = None
    for line in lines:
        has_school = bool(_SCHOOL.search(line))
        degree = _DEGREE.search(line)
        year = re.findall(r"(?:19|20)\d{2}", line)
        if has_school and (cur is None or cur.get("school")):
            cur = {"degree": "", "field": "", "school": "", "status": "completed", "year": ""}
            entries.append(cur)
        if cur is None and degree:
            cur = {"degree": "", "field": "", "school": "", "status": "completed", "year": ""}
            entries.append(cur)
        if cur is None:
            continue
        if has_school and not cur["school"]:
            cur["school"] = _clean_header(re.split(r"\s{2,}|\|", _DATE_RANGE.sub("", line))[0]) or _clean_header(line)
        if degree and not cur["degree"]:
            cur["degree"] = degree.group(0).strip()
            fm = _FIELD.search(line[degree.end():])
            if fm:
                cur["field"] = fm.group(1).strip(" ,")
        if year:
            cur["year"] = year[-1]
        if re.search(r"expected|present|current|in progress|candidate", line, re.IGNORECASE) or _ends_in_future(line):
            cur["status"] = "in_progress"
    return [e for e in entries if e["degree"] or e["school"]]


def _education_evidence(education: list[dict]) -> list[dict]:
    out = []
    for e in education:
        text = " ".join(x for x in [e["degree"], e["field"], "-", e["school"]] if x)
        out.append({"id": uuid.uuid4().hex[:10], "section": "Education", "text": text, "skills": [], "context": ""})
    return out


def parse_resume(text: str) -> dict:
    sections = split_sections(text)
    warnings: list[str] = []
    exp_items, _ = evidence_from_section(sections.get("experience", []), "Experience")
    proj_items, projects = evidence_from_section(sections.get("projects", []), "Projects")
    other_items, _ = evidence_from_section(sections.get("other", []), "Other")
    education = parse_education(sections.get("education", []))
    skills_text = "\n".join(sections.get("skills", []))
    listed = extract_skills(skills_text)
    all_text_skills = extract_skills(text)
    months = experience_months(sections.get("experience", []))
    if not sections.get("experience"):
        warnings.append("No Experience section detected; add evidence items manually.")
    elif months == 0:
        warnings.append("Could not detect dates in the Experience section; set experience months manually.")
    if not listed:
        warnings.append("No Skills section detected; skills were inferred from the whole document.")
    skills = listed + [s for s in all_text_skills if s not in listed]
    skill_evidence = [{
        "id": uuid.uuid4().hex[:10], "section": "Skills", "text": skills_text[:400],
        "skills": listed, "context": "Skills section",
    }] if listed else []
    return {
        "skills": skills,
        "technologies": listed or skills,
        "experience_months": months,
        "education": education,
        "evidence_items": exp_items + proj_items + other_items + skill_evidence + _education_evidence(education),
        "projects": [p for p in projects if p["text"]],
        "warnings": warnings,
    }
